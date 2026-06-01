"""
StudyMind — Từ vựng chuẩn JSON → Flashcard / Quiz
==================================================
Schema:
{
  "vocabulary": [
    {
      "tu_vung": "hello",
      "nghia": "xin chào",
      "vi_du": "Hello, how are you?",
      "phat_am": "/həˈloʊ/"
    }
  ]
}
"""

from __future__ import annotations

import csv
import io
import json
import os
import random
import re
from typing import Any

from dotenv import load_dotenv
from openai import AsyncOpenAI

load_dotenv()

_vocab_client = AsyncOpenAI(
    api_key=os.getenv("OPENAI_API_KEY"),
    base_url="https://openrouter.ai/api/v1",
)

_VOCAB_MODELS = [
    os.getenv("MODEL", "deepseek/deepseek-v4-flash:free"),
    "deepseek/deepseek-v4-flash:free",
    "meta-llama/llama-3.3-70b-instruct:free",
    "openai/gpt-oss-120b:free",
]

VOCABULARY_EXTRACT_PROMPT = """Bạn là chuyên gia trích xuất từ vựng học ngoại ngữ và học thuật.

NHIỆM VỤ: Đọc nội dung và trả về JSON từ vựng CHUẨN (không thêm text khác).

OUTPUT (chỉ JSON thuần):
{
  "vocabulary": [
    {
      "tu_vung": "từ hoặc cụm (ngôn ngữ gốc)",
      "nghia": "nghĩa tiếng Việt hoặc giải thích ngắn",
      "vi_du": "một câu ví dụ sử dụng từ",
      "phat_am": "phiên âm / IPA / romaji / pinyin nếu có"
    }
  ]
}

QUY TẮC:
- Mỗi mục phải có tu_vung và nghia
- Không bịa từ không có trong nội dung (có thể suy luận ví dụ hợp lý)
- Tối đa theo yêu cầu số lượng trong message
- Tiếng Hàn: tu_vung = Hangul, phat_am = romanization
- Tiếng Anh: phat_am = IPA hoặc phiên âm đơn giản"""


VOCAB_JSON_SCHEMA = {
    "vocabulary": [
        {
            "tu_vung": "string — từ/cụm cần học",
            "nghia": "string — nghĩa tiếng Việt hoặc giải thích",
            "vi_du": "string — câu ví dụ (tuỳ chọn)",
            "phat_am": "string — phiên âm / IPA / romanization (tuỳ chọn)",
        }
    ]
}

HEADER_ALIASES = {
    "tu_vung": ("tu_vung", "tu", "từ", "word", "term", "vocabulary", "mặt trước", "front", "kanji", "hangul"),
    "nghia": ("nghia", "nghĩa", "meaning", "definition", "dịch", "mặt sau", "back", "answer"),
    "vi_du": ("vi_du", "vi du", "ví dụ", "example", "cau_vi_du", "sentence"),
    "phat_am": ("phat_am", "phat am", "phát âm", "pronunciation", "ipa", "romanization", "romaji", "pinyin"),
}


def _norm_header(h: str) -> str:
    return re.sub(r"\s+", " ", (h or "").strip().lower())


def _map_columns(headers: list[str]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    norm = [_norm_header(h) for h in headers]
    for field, aliases in HEADER_ALIASES.items():
        for i, h in enumerate(norm):
            if h in aliases or any(a in h for a in aliases):
                mapped[field] = i
                break
    return mapped


def _row_to_vocab(row: list[str], col_map: dict[str, int]) -> dict | None:
    def cell(key: str) -> str:
        idx = col_map.get(key)
        if idx is None or idx >= len(row):
            return ""
        return (row[idx] or "").strip()

    tu = cell("tu_vung")
    nghia = cell("nghia")
    if not tu and not nghia:
        return None
    if not tu:
        tu = nghia
    if not nghia:
        nghia = tu
    return {
        "tu_vung": tu,
        "nghia": nghia,
        "vi_du": cell("vi_du"),
        "phat_am": cell("phat_am"),
    }


def normalize_vocab_item(item: dict) -> dict | None:
    if not item:
        return None
    tu = (
        item.get("tu_vung")
        or item.get("word")
        or item.get("term")
        or item.get("front")
        or item.get("question")
        or ""
    )
    nghia = (
        item.get("nghia")
        or item.get("meaning")
        or item.get("definition")
        or item.get("back")
        or item.get("answer")
        or ""
    )
    tu, nghia = str(tu).strip(), str(nghia).strip()
    if not tu and not nghia:
        return None
    if not tu:
        tu = nghia
    if not nghia:
        nghia = tu
    return {
        "tu_vung": tu,
        "nghia": nghia,
        "vi_du": str(item.get("vi_du") or item.get("example") or "").strip(),
        "phat_am": str(item.get("phat_am") or item.get("pronunciation") or "").strip(),
    }


def normalize_vocab_list(items: list) -> list[dict]:
    out = []
    seen = set()
    for raw in items or []:
        if not isinstance(raw, dict):
            continue
        v = normalize_vocab_item(raw)
        if not v:
            continue
        key = v["tu_vung"].lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(v)
    return out


def parse_paste_text(text: str) -> list[dict]:
    """
    Phân tích nội dung dán từ Excel, Word, JSON hoặc từng dòng.
    Hỗ trợ: Tab (copy Excel), |, dấu phẩy CSV, JSON, word - nghĩa, word: nghĩa
    """
    raw = (text or "").strip()
    if not raw:
        return []

    if raw.startswith("{") or raw.startswith("["):
        try:
            data = json.loads(raw)
            if isinstance(data, dict) and "vocabulary" in data:
                return normalize_vocab_list(data["vocabulary"])
            if isinstance(data, list):
                return normalize_vocab_list(data)
        except json.JSONDecodeError:
            pass

    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    if not lines:
        return []

    sample = lines[0]
    if "\t" in sample:
        rows = [[p.strip() for p in ln.split("\t")] for ln in lines]
        col_map = _map_columns(rows[0])
        start = 1 if col_map else 0
        if not col_map:
            col_map = {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3}
        parsed = [_row_to_vocab(r, col_map) for r in rows[start:]]
        out = [v for v in parsed if v]
        if out:
            return normalize_vocab_list(out)

    if "|" in sample:
        rows = [[p.strip() for p in ln.split("|")] for ln in lines]
        col_map = _map_columns(rows[0])
        start = 1 if col_map else 0
        if not col_map:
            col_map = {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3}
        parsed = [_row_to_vocab(r, col_map) for r in rows[start:]]
        out = [v for v in parsed if v]
        if out:
            return normalize_vocab_list(out)

    if "," in raw and raw.count(",") >= len(lines):
        csv_parsed = parse_vocabulary_csv(raw)
        if csv_parsed:
            return csv_parsed

    result = []
    col = {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3}
    for line in lines:
        if line.lower().startswith(("tu_vung", "từ vựng", "word", "từ\t")):
            continue
        matched = False
        for sep in (" - ", " – ", " — ", ":", ";", "|"):
            if sep in line:
                parts = [p.strip() for p in re.split(re.escape(sep), line, maxsplit=3)]
                v = _row_to_vocab(parts, col)
                if v:
                    result.append(v)
                matched = True
                break
        if not matched:
            parts = re.split(r"\s{2,}", line)
            if len(parts) >= 2:
                v = _row_to_vocab(parts, col)
                if v:
                    result.append(v)

    return normalize_vocab_list(result)


def parse_vocabulary_csv(text: str) -> list[dict]:
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    col_map = _map_columns(rows[0])
    data_rows = rows[1:] if col_map else rows
    if not col_map and rows:
        col_map = {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3}
    result = []
    for row in data_rows:
        v = _row_to_vocab(row, col_map)
        if v:
            result.append(v)
    return normalize_vocab_list(result)


def parse_vocabulary_xlsx(content: bytes) -> list[dict]:
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("Cần cài openpyxl: pip install openpyxl") from e

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([str(c).strip() if c is not None else "" for c in row])
    wb.close()
    if not rows:
        return []
    col_map = _map_columns(rows[0])
    data_rows = rows[1:] if col_map else rows
    if not col_map:
        col_map = {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3}
    result = []
    for row in data_rows:
        v = _row_to_vocab(row, col_map)
        if v:
            result.append(v)
    return normalize_vocab_list(result)


def parse_vocabulary_docx(content: bytes) -> list[dict]:
    from docx import Document as DocxDoc

    doc = DocxDoc(io.BytesIO(content))
    result: list[dict] = []

    for table in doc.tables:
        cells = table.rows
        if len(cells) < 2:
            continue
        headers = [c.text.strip() for c in cells[0].cells]
        col_map = _map_columns(headers)
        if not col_map:
            col_map = {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3}
        for row in cells[1:]:
            row_vals = [c.text.strip() for c in row.cells]
            v = _row_to_vocab(row_vals, col_map)
            if v:
                result.append(v)

    if result:
        return normalize_vocab_list(result)

    for para in doc.paragraphs:
        line = para.text.strip()
        if not line or line.startswith("#"):
            continue
        parts = re.split(r"[\t|,;]+", line)
        if len(parts) >= 2:
            v = _row_to_vocab(parts, {"tu_vung": 0, "nghia": 1, "vi_du": 2, "phat_am": 3})
            if v:
                result.append(v)

    return normalize_vocab_list(result)


def parse_vocabulary_file(content: bytes, filename: str) -> list[dict]:
    ext = (filename or "").rsplit(".", 1)[-1].lower()
    if ext == "csv" or ext == "txt":
        text = content.decode("utf-8", errors="ignore")
        return parse_vocabulary_csv(text)
    if ext in ("xlsx", "xls"):
        if ext == "xls":
            raise RuntimeError("File .xls cũ chưa hỗ trợ — lưu lại .xlsx hoặc .csv")
        return parse_vocabulary_xlsx(content)
    if ext == "docx":
        return parse_vocabulary_docx(content)
    if ext == "json":
        data = json.loads(content.decode("utf-8", errors="ignore"))
        items = data.get("vocabulary", data) if isinstance(data, dict) else data
        if not isinstance(items, list):
            raise RuntimeError("JSON phải có mảng vocabulary")
        return normalize_vocab_list(items)
    raise RuntimeError(f"Không hỗ trợ định dạng .{ext} — dùng .xlsx, .csv, .docx hoặc .json")


def extract_vocabulary_json(raw: str) -> list[dict]:
    cleaned = re.sub(r"```(?:json)?\s*", "", raw or "").replace("```", "").strip()
    data: Any = None
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if m:
            data = json.loads(m.group())
    if isinstance(data, dict) and "vocabulary" in data:
        return normalize_vocab_list(data["vocabulary"])
    if isinstance(data, list):
        return normalize_vocab_list(data)
    return []


def vocabulary_to_payload(vocab: list[dict]) -> dict:
    return {"vocabulary": normalize_vocab_list(vocab)}


def vocab_to_flashcards(vocab: list[dict]) -> list[dict]:
    cards = []
    for v in normalize_vocab_list(vocab):
        front = v["tu_vung"]
        if v["phat_am"]:
            front += f"\n({v['phat_am']})"
        back_parts = [v["nghia"]]
        if v["vi_du"]:
            back_parts.append(f"Ví dụ: {v['vi_du']}")
        cards.append({
            "question": front,
            "answer": "\n".join(back_parts),
            "hint": v["phat_am"] or "",
            "type": "vocabulary",
            "tu_vung": v["tu_vung"],
            "nghia": v["nghia"],
            "vi_du": v["vi_du"],
            "phat_am": v["phat_am"],
        })
    return cards


def vocab_to_quiz_questions(vocab: list[dict], num_questions: int | None = None) -> list[dict]:
    items = normalize_vocab_list(vocab)
    if len(items) < 2:
        return []
    random.shuffle(items)
    n = num_questions or len(items)
    n = min(n, len(items))
    questions = []

    for i in range(n):
        correct = items[i]
        others = [x for x in items if x["tu_vung"] != correct["tu_vung"]]
        random.shuffle(others)
        distractors = [o["nghia"] for o in others[:3]]
        while len(distractors) < 3 and len(others) > len(distractors):
            distractors.append(others[len(distractors)]["nghia"])
        while len(distractors) < 3:
            distractors.append(f"Nghĩa gây nhiễu {len(distractors) + 1}")

        options = [correct["nghia"]] + distractors[:3]
        random.shuffle(options)
        correct_index = options.index(correct["nghia"])

        q_text = f"Từ '{correct['tu_vung']}' có nghĩa là gì?"
        if correct["phat_am"]:
            q_text = f"Từ '{correct['tu_vung']}' ({correct['phat_am']}) có nghĩa là gì?"

        expl = f"Đáp án: {correct['nghia']}"
        if correct["vi_du"]:
            expl += f"\nVí dụ: {correct['vi_du']}"

        questions.append({
            "question": q_text,
            "options": options,
            "correctIndex": correct_index,
            "correct_index": correct_index,
            "explanation": expl,
            "tu_vung": correct["tu_vung"],
        })
    return questions


async def ai_extract_vocabulary(
    text: str,
    topic: str = "",
    max_items: int = 30,
) -> list[dict]:
    """Trích xuất từ vựng bằng LLM → danh sách chuẩn hoá."""
    sample = text[:6000]
    prompt = (
        f"Chủ đề: {topic or 'từ vựng'}\n"
        f"Trích xuất tối đa {max_items} mục từ vựng từ nội dung sau:\n\n{sample}"
    )
    for model in _VOCAB_MODELS:
        try:
            res = await _vocab_client.chat.completions.create(
                model=model,
                max_tokens=4000,
                messages=[
                    {"role": "system", "content": VOCABULARY_EXTRACT_PROMPT},
                    {"role": "user", "content": prompt},
                ],
            )
            raw = res.choices[0].message.content or ""
            items = extract_vocabulary_json(raw)
            if items:
                return items[:max_items]
        except Exception:
            continue
    return []
